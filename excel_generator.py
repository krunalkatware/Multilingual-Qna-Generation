import io
import openpyxl
from openpyxl.styles import Font, Alignment


def create_excel_workbook(english_qna, hindi_qna, marathi_qna):
    wb = openpyxl.Workbook()

    ws_english = wb.active
    ws_english.title = "English"
    ws_hindi = wb.create_sheet("Hindi")
    ws_marathi = wb.create_sheet("Marathi")

    sheets = [
        (ws_english, english_qna),
        (ws_hindi, hindi_qna),
        (ws_marathi, marathi_qna),
    ]

    header_font = Font(name="Calibri", size=11, bold=True)
    cell_alignment = Alignment(wrap_text=True, vertical="top")

    for ws, data in sheets:
        ws.cell(row=1, column=1, value="Questions").font = header_font
        ws.cell(row=1, column=2, value="Answers").font = header_font
        ws.freeze_panes = "A2"

        for row_idx, item in enumerate(data, start=2):
            q_cell = ws.cell(row=row_idx, column=1, value=item.get("question", ""))
            a_cell = ws.cell(row=row_idx, column=2, value=item.get("answer", ""))
            q_cell.alignment = cell_alignment
            a_cell.alignment = cell_alignment

        ws.column_dimensions["A"].width = 45
        ws.column_dimensions["B"].width = 75

    return wb


def generate_excel_bytes(english_qna, hindi_qna, marathi_qna):
    wb = create_excel_workbook(english_qna, hindi_qna, marathi_qna)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    data = output.getvalue()

    is_valid, error = validate_excel(data)
    if not is_valid:
        raise ValueError(f"Excel validation failed: {error}")

    return data


def save_excel_file(english_qna, hindi_qna, marathi_qna, file_path="Multilingual_QnA.xlsx"):
    wb = create_excel_workbook(english_qna, hindi_qna, marathi_qna)
    wb.save(file_path)

    is_valid, error = validate_excel(file_path)
    if not is_valid:
        raise ValueError(f"Excel validation failed: {error}")

    return file_path


def validate_excel(file_or_bytes):
    try:
        if isinstance(file_or_bytes, bytes):
            wb = openpyxl.load_workbook(io.BytesIO(file_or_bytes))
        else:
            wb = openpyxl.load_workbook(file_or_bytes)

        expected_sheets = ["English", "Hindi", "Marathi"]
        if wb.sheetnames != expected_sheets:
            return False, f"Expected sheets {expected_sheets}, got {wb.sheetnames}"

        row_counts = []
        for name in expected_sheets:
            ws = wb[name]
            col1 = ws.cell(row=1, column=1).value
            col2 = ws.cell(row=1, column=2).value
            col3 = ws.cell(row=1, column=3).value

            if col1 != "Questions" or col2 != "Answers" or col3 is not None:
                return False, f"Sheet '{name}' columns must be exactly 'Questions' and 'Answers'."

            count = ws.max_row - 1
            row_counts.append(count)

        if not (row_counts[0] == row_counts[1] == row_counts[2]):
            return False, f"Row count mismatch across sheets: {row_counts}"

        return True, None
    except Exception as e:
        return False, str(e)
